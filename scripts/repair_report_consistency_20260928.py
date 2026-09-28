"""Repair the reviewed v3 manuscript without refitting models or altering originals.

Run with the bundled document runtime. Numeric appendix rows come from frozen CSVs.
Only document.xml and the verified replacement plot are changed in the DOCX package.
"""
from pathlib import Path
from copy import deepcopy
import argparse, csv, difflib, hashlib, json, re, zipfile
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def readcsv(p):
    with p.open(encoding='utf-8-sig', newline='') as f: return list(csv.DictReader(f))
def number(x): return f'{float(x):.6f}'.replace('-', '−')

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--input',type=Path,required=True); ap.add_argument('--output',type=Path,required=True); ap.add_argument('--data-root',type=Path,required=True); ap.add_argument('--results',type=Path,required=True); a=ap.parse_args()
    assert a.input != a.output
    if a.output.exists():
        prior=json.loads((a.results/'audit15_edit_manifest.json').read_text(encoding='utf-8'))
        assert str(a.output)==prior['output'] and sha(a.output)==prior['output_sha256'], 'Refuse to overwrite an unrecognized file'
    d=Document(a.input); ps=list(d.paragraphs); assert len(ps)==751 and len(d.tables)==51
    changes=[]
    def settext(p,text,label,all_yellow=False):
        old=p.text
        if old==text and not all_yellow: return
        spans=[]; start=0
        for run in p.runs:
            spans.append((start,start+len(run.text),deepcopy(run._r.rPr))); start+=len(run.text)
        default=next((x[2] for x in spans if x[2] is not None),None)
        def styleat(i): return next((st for lo,hi,st in spans if lo<=i<hi),default)
        for el in list(p._p):
            if el.tag!=qn('w:pPr'): p._p.remove(el)
        def put(t,st,yellow):
            if not t: return
            run=p.add_run(t)
            if st is not None: run._r.insert(0,deepcopy(st))
            if yellow:
                pr=run._r.get_or_add_rPr(); h=pr.find(qn('w:highlight'))
                if h is None: h=OxmlElement('w:highlight'); pr.append(h)
                h.set(qn('w:val'),'yellow')
        if all_yellow: put(text,default,True)
        else:
            for op,i,j,k,l in difflib.SequenceMatcher(None,old,text,autojunk=False).get_opcodes():
                if op=='equal':
                    for lo,hi,st in spans:
                        if min(j,hi)>max(i,lo): put(old[max(i,lo):min(j,hi)],st,False)
                elif op in ('insert','replace'): put(text[k:l],styleat(i),True)
        changes.append(dict(location=label,before=old,after=text))
    def p(i,text): settext(ps[i],text,f'paragraph {i}')
    def addafter(i,text):
        el=deepcopy(ps[i]._p); ps[i]._p.addnext(el); pp=Paragraph(el,d); settext(pp,text,f'insert after {i}',True); return pp
    replacements={
      147: ps[147].text.replace('6.25%','5.6%').replace('三个条件中的 No-Go 比例分别为 22.2%、11.1% 和 5.6%','三个条件中的 No-Go 次数分别为 48、24 和 12 次，占 216 个试次的 22.2%、11.1% 和 5.6%'),
      164: '注：图示为第二阶段预实验 19 份有效记录在三个 B 区块中的描述性轨迹，灰线表示单份记录，深线表示记录均值。阴影为原分析按记录计算的 95% 置信区间，未按重复参与者聚类；本图仅用于形成性描述，不作为 19 名独立参与者的推断证据。',
      184: ps[184].text.replace('时间窗口至少包含 20 个有效正确 Go RT 才计算 RT-CV。','每个分析窗口至少包含 2 个不低于 100 ms 的有效正确 Go RT，且均值为正时，才以样本标准差除以均值计算 RT-CV。2 次为数学可计算条件，不代表该窗口已经具有充分的测量精度。'),
      188: '注：本图来自第二阶段预实验的历史行为分析，展示 19 份有效记录中的 20941 次正确 Go 反应，包括直方图与经验累积分布。100、150、1000 和 1150 ms 仅作质量检查标记，图中未据此删除反应。本图用于说明分布与检查阈值，不表示正式 116 场的试次总量；正式行为指标采用正文所述 100 ms 条件。',
      191: 'Go 遗漏进一步依据按键时序特征区分为无已检测时序歧义遗漏和时序歧义遗漏。前者指未检出预设异常按键时序特征的 Go 遗漏，后者指伴有刺激前按键或跨试次按键延续等特征的 Go 遗漏。两类遗漏率均以全部有效 Go 试次为分母，其和等于原始 Go 遗漏率。该划分描述记录中的时序特征，不直接判定参与者是否真正脱离任务或有意预判；迟发反应另作候选时序指标。',
      195: '近红外视频首先使用项目训练的 YOLO26n 检测器定位眼部区域（Ultralytics, 2026），再采用 RITnet 四分类语义分割模型识别背景、巩膜、虹膜和瞳孔（Chaudhary et al., 2019）。本报告主要眼部关联结果及预测模型采用硬分割瞳孔面积比例，并由此形成水平、波动、线性趋势和二次曲率四项瞳孔指标。几何平均直径由拟合椭圆的长短轴计算，单位为 px，保留为替代表示与测量敏感性分析。两种表示的测量尺度不同，其比较结果见 5.3.1。',
      207: '身体运动既提供任务中的外显行为信息，也会影响微弱心肺信号的估计。毫米波分析分别提取心率、呼吸频率、运动代理和质量指标。当前测量结果报告有效覆盖、内部质量分布及运动代理与质量指标的描述性相关，心率和呼吸率仅作支持性估计。逐搏间期与心率变异性的应用留待独立同步参考验证。',
      208: '逐帧毫米波数据按照采集过程中记录的实际时间与任务事件对应。开发阶段采用区块内事件标记与心电采样点的仿射映射检查同步，不将单次对齐残差推广为全部场次的同步精度。预测输入采用后续完成时间合同核验的生产版本，按探针前时间范围选取原始帧；旧测量快照及开发估计保留各自的时间来源和用途。',
      209: ps[209].text.replace('滤波阶数、VMD 参数和内部搜索窗口等实现细节列于附录 A。','当前生产实现采用四阶巴特沃斯带通及双向滤波；变分模态分解设定为惩罚参数 1000、模态数 3、更新步长 0、非固定直流分量、均匀初始化和收敛阈值 10⁻⁶。具体生产版本与开发版本的用途见附录 E.6。'),
      213: '注：本图为历史信号处理示例。A 为平均距离谱与选中的距离单元；B 展示 20 s 的相位解调片段及呼吸（0.10～0.50 Hz）和心搏（0.8～2.0 Hz）带通成分，后两者为显示清晰作纵向偏移；C 标出的频谱峰对应 55 次/分。图中的示例片段和谱峰不代表正式预测的 30 s 窗口或融合心率估计。',
      217: ps[217].text.replace('线性趋势和、二次曲率','线性趋势和二次曲率'),
      238: '信息类别层评价行为、眼部、动作和心肺估计对注意内容报告的预测能力，并比较在行为基础上加入各类传感信息的变化。眼部、动作与心肺类别模型分别采用其与行为五项均有效的共同比较集合，并非各类信息自身的全部有效记录。独立单指标和八项传感联合按各自所需输入形成集合，不附加行为完整性条件。',
      275: '注：A 为同一场次中 B1 与 B2 的配对轨迹和场次均值；B 为各区块内六个连续任务阶段的轨迹，先在参与者内汇总，再计算参与者均值与标准误。阶段表示每区块 24 个刺激循环合并形成的六段。图中误差线用于描述，正式推断以参与者聚类分析为准。',
      317: '注：图中点为累积优势比，横线为 95% 置信区间，虚线为无关联值 1；优势比大于 1 对应较高清醒等级。正文报告对数优势系数，图中数值为其指数变换。q 值在参与者内外分量共 10 项检验中校正。',
      330: ps[330].text.replace(' s⁻¹','')+' 动作强度为无量纲相对指标，阶段系数表示每推进一个连续阶段的变化量。',
      344: '2026 年 9 月 12 日的毫米波测量评估快照中，109 个场次的 2180 个探针窗口同时获得可计算的心率和呼吸率估计，占全部 2320 个探针的 93.97%。其余 40 个探针对应来源不可用的 2 场，100 个探针对应来源格式损坏的 5 场。现有审计不足以将后者统一归因为设备休眠。',
      347: '旧测量快照中，2320 条输出记录的窗口终点均与探针时间一致，该项检查仅确认输出端点。后续生产版本另完成原始帧选择、区块边界及探针前时间支持的合同核验，用于预测的心率和呼吸率各有 2200 个有限值探针，行为与心肺共同比较保留 2198 个。旧快照的 2180 个窗口继续用于本节描述统计，三者分别对应不同版本或筛选条件。',
      352: '5.5.3 开发场次的同步参考对照',
      353: '同步心电图（electrocardiogram [ECG]）与呼吸带对照来自同一参与者的重复开发场次，其中 5 个可估计场次提供 100 个探针窗口。2026 年 8 月 31 日的估计中，30 s 融合心率平均绝对误差为 10.83 次/分；末尾 25 s 估计仅有 65 个可估计窗口，平均绝对误差为 9.12 次/分，中位绝对误差为 5.79 次/分；呼吸频率平均绝对误差为 4.17 次/分。25 s 与 30 s 的有效集合不同，不能据此直接认定较短窗口更准确。后续切换帧时间来源后的 100 窗口冻结开发诊断中，融合心率平均绝对误差为 10.457 次/分，平均符号误差为 −9.033 次/分，详见附录 E。以上均用于估计器开发，不构成独立参与者的生理测量验证。',
      356: '注：本图保留早期 3 个开发场次、60 个窗口的历史对照，来自同一参与者的重复记录。图示偏倚为 −6.4 次/分，一致性界限为 −28.6 至 15.8 次/分，仅描述该历史子集，不代表后续 5 场、100 窗口的冻结结果，也不作为参与者总体的一致性界限。',
      362: ps[362].text.replace('心肺类别采用','眼部与动作类别分别采用行为共同集合的 1775 和 2296 个探针；心肺类别采用'),
      367: ps[367].text.replace('各行按自身集合解释。','眼部、动作和心肺类别采用与行为的共同集合，其余行按所列集合解释。'),
      371: ps[371].text.replace('各类输入分别采用自身可用场次。','各行沿用表 14 的分析集合，再保留可计算该指标的场次。'),
      383: '5.6.4 行为窗口长度的补充结果',
      384: ps[384].text.replace('表 5-13','表 16'),
      388: '5.6.5 四分类预测与补充分析',
      405: ps[405].text+' 在行为基础上加入心肺估计的改善量为 0.00050，95% 置信区间为 [−0.00576, 0.00692]，采用 2198 个共同探针。',
      407: '注：每行均沿用该比较共同样本上完成的行为基线与增量模型训练结果，心肺比较采用 58 人、110 场、2198 个探针。各项改善量按其共同样本解释。',
      408: ps[408].text.replace('线性概率模型','L2 正则化逻辑回归'),
      424: ps[424].text+' 心率和呼吸率独立单指标均采用 58 人、110 场、2200 个有效探针，损失分别为 0.6331 [0.5814, 0.6840] 和 0.6346 [0.5858, 0.6843]。在 2198 个行为共同探针上，单独加入心率与呼吸率的改善量分别为 0.00218 [−0.00361, 0.00836] 和 −0.00166 [−0.00284, −0.00062]。',
      445: '注：点为优势比，横线为 95% 置信区间，虚线为 1。橙色表示确认性效标一致性模型，灰色表示探索性模型，颜色不表示单项区间是否排除 1。No-Go 误按率的单项区间排除 1，但经探索性错误发现率校正后 q = .143。',
      451: ps[451].text+' 本节展示交互原型及示例数据。0–100 状态指数、专注时长、稳定片段和画像均为拟议的呈现方式，尚未建立与已验证预测输出的对应关系；图中若出现心率变异性，也仅为待验证的预留字段。',
      456: ps[456].text.replace('实线框表示已完成并获得相应研究支持的环节，虚线框表示已实现的界面原型功能，其测量应用随后续模型验证逐步完善。','实线框表示测量研究流程，虚线框表示界面原型；各环节的验证程度以本报告对应结果为准，连线不表示原型指标已经获得测量验证。'),
      458: ps[458].text+' 当前二分类模型估计随后报告任务聚焦的概率，该概率尚不能直接换算为注意强度或专注时间占比；原型中的可信程度也不等同于已完成校准的个体预测置信度。',
      498: '注：界面原型集中呈现设备与检查、模型范围、人工智能使用声明与数据边界四项说明。版本与适用场景是拟议的信息字段，图中示例状态不代表已经发布并完成应用验证的注意测量结果。',
      626: '· 本节场次级问卷模型沿用其历史预设条件，以不低于 150 ms 的正确 Go 反应计算中位数；该条件与正式探针窗口行为表示的不低于 100 ms 分别对应不同分析，未相互替换。',
      691: '注：采用以参与者为簇、交换型工作相关结构的广义估计方程（generalized estimating equations [GEE]）。系数保留动作指标原尺度；区块、阶段及其交互项未按一个标准差解释。整体动作强度和前后方向为无量纲，横向与纵向方向为归一化方向分数/秒。各模型均为 61 人、115 场、1380 条阶段记录。95% 置信区间对应单项估计。',
      694: '注：采用参与者聚类稳健多项逻辑回归，以任务聚焦为参照。2、3、4 分别表示关注实验但未聚焦分类任务、任务无关思维与思维空白。系数为动作预测变量增加一个标准差时的对数优势变化，均为 61 人、115 场、2300 个探针；区间未作整组多重校正。',
      697: '注：采用以参与者为簇的有序广义估计方程，因变量为四级困倦—清醒报告。系数为动作预测变量增加一个标准差时的累积对数优势变化；正值对应较高清醒等级。各模型均为 61 人、115 场、2300 个探针。',
      700: '注：采用以参与者为簇的高斯广义估计方程。因变量为所列动作指标，预测变量为同期行为指标；系数表示行为指标增加一个标准差时，动作指标原尺度的变化量。反应时变异系数在每个探针前窗口内，按至少 2 次有效正确 Go 反应计算，并非每名参与者累计 2 次。各模型含 61 人、115 场，实际窗口数按行报告。',
      703: '注：有效率为有限值探针数除以全部 2320 个探针。整体动作强度为主要候选；横向、纵向方向为敏感性辅助指标，前后方向仅作质量控制与敏感性描述；曝光变化、全画面运动和平均灰度仅作设备质量控制。覆盖数不等同于某次行为共同模型的样本数。',
      709: ps[709].text+' 此项为旧快照的输出端点检查；后续预测生产版本的原始帧时间合同核验及不同覆盖口径见 5.5.1。',
      711: ps[711].text.replace('在 5 个同时具有合格心电参考的场次','在同一参与者的 5 个具有合格心电参考的重复开发场次'),
      725: ps[725].text.replace('基于 5 个场次','基于同一参与者的 5 个重复开发场次'),
    }
    replacements[195]=replacements[195].replace('采用硬分割瞳孔面积比例，并由此','采用硬分割瞳孔面积占瞳孔与虹膜合计面积的比例，并由此')
    replacements[330]=ps[330].text+' 动作强度为归一化帧间灰度差除以帧间隔所得的变化率，单位为 s⁻¹；上述系数表示每推进一个连续任务阶段，该变化率的变化量，不表示物理位移或能量。'
    replacements[691]=replacements[691].replace('整体动作强度和前后方向为无量纲，横向与纵向方向为归一化方向分数/秒。','整体动作强度为归一化画面变化率，单位为 s⁻¹；横向与纵向方向为归一化方向分数/秒，前后方向分数无量纲。')
    replacements.update({177:ps[177].text.replace('表 2','表 7'),267:ps[267].text.replace('表 5-3','表 10'),286:ps[286].text.replace('表 5-5','附录 C.2'),300:ps[300].text.replace('图 20','图 21')})
    for i,t in replacements.items(): p(i,t)
    addafter(137,'正式主分析的 116 个场次均属于北京采集批次，每场包含两个 B 区块和 20 次思维探针，共 2320 次。珠海批次保留其不同版本的探针记录，未并入本报告的正式主队列；两阶段预实验也不并入该分母。')
    # Retire the redundant cardiopulmonary subsection after redistributing its results.
    for i in (380,381,382):
        changes.append(dict(location=f'delete paragraph {i}',before=ps[i].text,after='内容并入 5.6.1、5.7.1、5.7.3 和 5.8'))
        ps[i]._p.getparent().remove(ps[i]._p)
    # Targeted old references and terminology, including tables.
    allp=list(d.paragraphs)+[p for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs]
    for pp in allp:
        text=pp.text
        text=text.replace('5.6.6','5.6.5').replace('线性概率模型','L2 正则化逻辑回归').replace('心肺（补充）','心肺').replace('表 5-3','表 10').replace('表 5-5','附录 C.2')
        text=text.replace('真遗漏','无已检测时序歧义遗漏').replace('预判遗漏','时序歧义遗漏').replace('明确遗漏','无已检测时序歧义遗漏')
        text=text.replace('Go 遗漏率（清洗后）','无已检测时序歧义遗漏率').replace('行为结局','行为指标').replace('结局分布','结果变量分布').replace('结局为','因变量为').replace('结局','结果变量')
        if pp is ps[177]: text=text.replace('表 2','表 7')
        if pp is ps[267]: text=text.replace('表 5-3','表 10')
        if pp is ps[286]: text=text.replace('表 5-5','附录 C.2')
        if pp is ps[300]: text=text.replace('图 20','图 21')
        if text!=pp.text: settext(pp,text,'term/reference '+pp.text[:28])
    # Table C.1 repeats raw omission under an alias; preserve one source estimate.
    tc=d.tables[34]
    seen_raw=False
    for row in list(tc.rows)[1:]:
        if row.cells[0].text=='Go 遗漏率（原始）':
            if not seen_raw:
                seen_raw=True
                continue
            changes.append(dict(location='C.1 duplicate raw omission',before=[c.text for c in row.cells],after='保留同表 Go 遗漏率唯一行'))
            row._tr.getparent().remove(row._tr)
    labels={'body_motion_energy_median':'整体动作强度','pose_lateral_right_per_sec_median':'横向方向','pose_vertical_up_per_sec_median':'纵向方向','pose_radial_proximity_direction_score_median':'前后方向','exposure_change_abs_median':'曝光变化','global_motion_energy_median':'全画面运动','gray_mean_median':'平均灰度','go_correct_rt_median_ms':'正确 Go 反应时中位数','go_correct_rt_cv':'反应时变异系数','raw_go_omission_rate':'原始 Go 遗漏率','omission_rate':'原始 Go 遗漏率','commission_rate':'No-Go 误按率','dprime_loglinear':'辨别力 d′'}
    terms={'Intercept':'截距','block2':'区块 2','cycle_bin':'阶段','block2:cycle_bin':'区块 2×阶段'}
    root=a.data_root/'FormalScience/Movement/tables'
    def ci(x): return '['+number(x['ci_low'])+', '+number(x['ci_high'])+']'
    def table(idx,rows,widths):
        t=d.tables[idx]; templates=[deepcopy(t.rows[0]._tr),deepcopy(t.rows[1]._tr)]
        for el in list(t._tbl):
            if el.tag==qn('w:tr'):t._tbl.remove(el)
        grid=t._tbl.tblGrid
        for el in list(grid):grid.remove(el)
        for width in widths:
            el=OxmlElement('w:gridCol');el.set(qn('w:w'),str(width));grid.append(el)
        for ri,values in enumerate(rows):
            row=deepcopy(templates[0 if ri==0 else 1]);cells=row.findall(qn('w:tc'));tmpl=deepcopy(cells[0])
            for c in cells:row.remove(c)
            for value,width in zip(values,widths):
                c=deepcopy(tmpl);row.append(c);pr=c.find(qn('w:tcPr'))
                if pr is None:pr=OxmlElement('w:tcPr');c.insert(0,pr)
                for bad in list(pr):
                    if bad.tag in [qn('w:gridSpan'),qn('w:vMerge')]:pr.remove(bad)
                w=pr.find(qn('w:tcW'))
                if w is None:w=OxmlElement('w:tcW');pr.append(w)
                w.set(qn('w:w'),str(width));w.set(qn('w:type'),'dxa')
                pp=c.find(qn('w:p'))
                for extra in list(c):
                    if extra.tag==qn('w:p') and extra is not pp:c.remove(extra)
                settext(Paragraph(pp,d),str(value),f'table {idx} row {ri}',True)
            if ri==0:
                pr=row.find(qn('w:trPr'))
                if pr is None:pr=OxmlElement('w:trPr');row.insert(0,pr)
                if pr.find(qn('w:tblHeader')) is None:pr.append(OxmlElement('w:tblHeader'))
            t._tbl.append(row)
    sources={k:readcsv(root/(k+'.csv')) for k in ['movement_task_progression','movement_q1_models','movement_q2_models','movement_behavior_links','movement_feature_coverage']}
    table(37,[['动作指标','模型项','系数','95% 置信区间','阶段数']]+[[labels[x['metric']],terms[x['term']],number(x['estimate']),ci(x),x['n_rows']] for x in sources['movement_task_progression']],[1500,1550,1150,2850,900])
    table(38,[['动作指标','对比类别','系数','95% 置信区间','探针数']]+[[labels[x['predictor']],x['contrast_category'],number(x['estimate_per_predictor_sd']),ci(x),x['n_rows']] for x in sources['movement_q1_models']],[1500,1550,1150,2850,900])
    table(39,[['动作指标','系数','95% 置信区间','探针数']]+[[labels[x['predictor']],number(x['estimate_per_predictor_sd']),ci(x),x['n_rows']] for x in sources['movement_q2_models']],[2200,1700,3000,1050])
    table(40,[['动作因变量','行为预测变量','系数','95% 置信区间','窗口数']]+[[labels[x['outcome']],labels[x['predictor']],number(x['estimate_per_predictor_sd']),ci(x),x['n_rows']] for x in sources['movement_behavior_links']],[1400,1650,1150,2850,900])
    table(41,[['指标','有限值探针数','有效率','参与者数','场次数']]+[[labels[x['predictor_column']],x['finite_probe_n'],f"{float(x['finite_fraction'])*100:.2f}%",x['participant_group_n'],x['session_n']] for x in sources['movement_feature_coverage']],[2500,1700,1450,1300,1000])
    # Reproducibility detail belongs to the heart appendix rather than questionnaire Appendix A.
    h=addafter(726,'E.6 生产版本与开发估计的对应'); h.style=ps[722].style
    for run in h.runs: run.bold=True
    el=deepcopy(ps[725]._p); h._p.addnext(el)
    settext(Paragraph(el,d),'测量描述采用 2026 年 9 月 12 日冻结快照的 2180 个窗口；预测采用后续通过时间合同核验的生产输入，两个心肺指标各有 2200 个有限值，行为共同比较为 2198 个。开发误差采用同一参与者的 5 场、100 窗口诊断。生产实现的四阶带通、双向滤波和三模态变分分解参数见 4.5.4；代码入口为毫米波正式仓库 scripts/process_vital_signs_v3_1_1.py，分解后端固定为 sktime 1.1.0。30 s 探针窗口与开发阶段末尾 25 s 子窗口分别保留，历史信号示例不用于推断预测输入的时间范围。','E.6 content',True)
    refs=[
      'Bazarevsky, V., Grishchenko, I., Raveendran, K., Zhu, T., Zhang, F., & Grundmann, M. (2020). BlazePose: On-device real-time body pose tracking. arXiv. https://doi.org/10.48550/arXiv.2006.10204',
      'Chaudhary, A. K., Kothari, R., Acharya, M., Dangi, S., Nair, N., Bailey, R., Kanan, C., Diaz, G., & Pelz, J. B. (2019). RITnet: Real-time semantic segmentation of the eye for gaze tracking. In 2019 IEEE/CVF International Conference on Computer Vision Workshop (ICCVW) (pp. 3698–3702). IEEE. https://doi.org/10.1109/ICCVW.2019.00568',
      'Dragomiretskiy, K., & Zosso, D. (2014). Variational mode decomposition. IEEE Transactions on Signal Processing, 62(3), 531–544. https://doi.org/10.1109/TSP.2013.2288675',
      'Efron, B., & Tibshirani, R. J. (1994). An introduction to the bootstrap. Chapman & Hall. https://doi.org/10.1201/9780429246593',
      'Hautus, M. J. (1995). Corrections for extreme proportions and their biasing effects on estimated values of d′. Behavior Research Methods, Instruments, & Computers, 27(1), 46–51. https://doi.org/10.3758/BF03203619',
      'Kartynnik, Y., Ablavatski, A., Grishchenko, I., & Grundmann, M. (2019). Real-time facial surface geometry from monocular video on mobile GPUs. arXiv. https://doi.org/10.48550/arXiv.1907.06724',
      'Lugaresi, C., Tang, J., Nash, H., McClanahan, C., Uboweja, E., Hays, M., Zhang, F., Chang, C.-L., Yong, M. G., Lee, J., Chang, W.-T., Hua, W., Georg, M., & Grundmann, M. (2019). MediaPipe: A framework for building perception pipelines. arXiv. https://doi.org/10.48550/arXiv.1906.08172',
      'Macmillan, N. A., & Creelman, C. D. (2005). Detection theory: A user’s guide (2nd ed.). Lawrence Erlbaum Associates. https://doi.org/10.4324/9781410611147',
      'Sen, P. K. (1968). Estimates of the regression coefficient based on Kendall’s tau. Journal of the American Statistical Association, 63(324), 1379–1389. https://doi.org/10.1080/01621459.1968.10480934',
      'Soukupová, T., & Čech, J. (2016). Real-time eye blink detection using facial landmarks. In Proceedings of the 21st Computer Vision Winter Workshop. https://cmp.felk.cvut.cz/ftp/articles/cech/Soukupova-CVWW-2016.pdf',
      'Theil, H. (1950). A rank-invariant method of linear and polynomial regression analysis, I–II. Indagationes Mathematicae, 12, 85–91, 173–177. https://ir.cwi.nl/pub/18445',
      'Ultralytics. (2026). Ultralytics YOLO26 [Computer software documentation]. https://docs.ultralytics.com/models/yolo26/',
    ]
    existing=[p for p in ps[535:610] if re.match(r'^[A-Za-z].*\(\d{4}',p.text)]
    italic_titles={
        'Bazarevsky': 'BlazePose: On-device real-time body pose tracking',
        'Chaudhary': '2019 IEEE/CVF International Conference on Computer Vision Workshop (ICCVW)',
        'Dragomiretskiy': 'IEEE Transactions on Signal Processing, 62',
        'Efron': 'An introduction to the bootstrap',
        'Hautus': 'Behavior Research Methods, Instruments, & Computers, 27',
        'Kartynnik': 'Real-time facial surface geometry from monocular video on mobile GPUs',
        'Lugaresi': 'MediaPipe: A framework for building perception pipelines',
        'Macmillan': 'Detection theory: A user’s guide',
        'Sen': 'Journal of the American Statistical Association, 63',
        'Soukupová': 'Proceedings of the 21st Computer Vision Winter Workshop',
        'Theil': 'Indagationes Mathematicae, 12',
        'Ultralytics': 'Ultralytics YOLO26',
    }
    for ref in sorted(refs):
        target=next((p for p in existing if p.text.casefold()>ref.casefold()),ps[610])
        el=deepcopy(existing[0]._p);target._p.addprevious(el); pp=Paragraph(el,d);settext(pp,ref,'reference '+ref.split('(')[0],True)
        fmt=deepcopy(pp.runs[0]._r.rPr)
        title=next(value for key,value in italic_titles.items() if ref.startswith(key))
        before,after=ref.split(title,1)
        for run in list(pp.runs): pp._p.remove(run._r)
        for text,is_italic in [(before,False),(title,True),(after,False)]:
            run=pp.add_run(text);run._r.insert(0,deepcopy(fmt));run.italic=is_italic
    # Crop only obsolete raster headings; plots themselves remain byte-identical.
    for idx,top in [(162,3600),(186,6500)]:
        fill=ps[idx]._p.xpath('.//pic:blipFill')[0];sr=fill.find(qn('a:srcRect'))
        if sr is None: sr=OxmlElement('a:srcRect');fill.insert(1,sr)
        sr.set('t',str(top)); changes.append(dict(location=f'figure crop paragraph {idx}',before='历史图内标题',after='改用标黄图注界定样本；数据图保持原始字节'))
    # Replace Figure 17 using the frozen aggregate tables, with stage and SEM labels.
    part=d.part.related_parts[ps[273]._p.xpath('.//a:blip')[0].get(qn('r:embed'))]
    media={str(part.partname).lstrip('/'):(a.results/'audit15_fig17.png').read_bytes()}
    p(274,ps[274].text) # Figure caption remains; the changed explanatory note is highlighted.
    # Some inherited captions are literal text and others are SEQ fields. Word export
    # would otherwise renumber the remaining fields as if literal captions did not exist.
    for pp in d.paragraphs:
        if not re.match(r'^[图表]\s*\d',pp.text): continue
        fields=pp._p.xpath('.//w:fldChar[@w:fldCharType="begin"]')
        simple=pp._p.xpath('.//w:fldSimple')
        if fields or simple:
            for field in fields+simple: field.set(qn('w:fldLock'),'true')
            changes.append(dict(location='caption numbering lock',before=pp.text,after=pp.text))
            for run in pp.runs:
                if re.fullmatch(r'\s*\d+\s*',run.text):
                    pr=run._r.get_or_add_rPr();highlight=pr.find(qn('w:highlight'))
                    if highlight is None:highlight=OxmlElement('w:highlight');pr.append(highlight)
                    highlight.set(qn('w:val'),'yellow')
    xml=etree.tostring(d._element,xml_declaration=True,encoding='UTF-8',standalone=True)
    with zipfile.ZipFile(a.input) as zin,zipfile.ZipFile(a.output,'w',zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():zout.writestr(item,xml if item.filename=='word/document.xml' else media.get(item.filename,zin.read(item.filename)))
    out=Document(a.output);assert len(out.tables)==51
    manifest=dict(input=str(a.input),input_sha256=sha(a.input),output=str(a.output),output_sha256=sha(a.output),changes=changes,appendix_sources={k:sha(root/(k+'.csv')) for k in sources},replacement_media=list(media),new_references=refs,models_refitted=0)
    (a.results/'audit15_edit_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'changed_regions':len(changes),'output_sha256':manifest['output_sha256'],'models_refitted':0},ensure_ascii=False))

if __name__=='__main__':main()
