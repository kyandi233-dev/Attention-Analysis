"""Synchronize the verified report copy and aggregate evidence into Formal indices."""
from pathlib import Path
import argparse,json,shutil,re
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--docx',type=Path,required=True);ap.add_argument('--formal-repo',type=Path,required=True);ap.add_argument('--results',type=Path,required=True);a=ap.parse_args();r=a.formal_repo;d=Document(a.docx)
    paras=list(d.paragraphs);index={p._p:i for i,p in enumerate(paras)}
    figmap={261:'fig16',366:'fig28',378:'fig29',410:'fig32',419:'fig33',431:'fig34'}
    for name in figmap.values():shutil.copy2(a.results/'figures'/(name+'.png'),r/'国赛报告/assets'/('20260928_'+name+'.png'))
    def mdtable(t):
        lines=[]
        for i,row in enumerate(t.rows):
            lines.append('| '+' | '.join(c.text.replace('\n','<br>').replace('|','\\|') for c in row.cells)+' |')
            if i==0:lines.append('| '+' | '.join('---' for c in row.cells)+' |')
        return '\n'.join(lines)
    def extract(start,end):
        active=False;result=[]
        for el in d._element.body:
            if el.tag==qn('w:p'):
                i=index.get(el,-1)
                if i==start:active=True
                if i==end:break
                if not active:continue
                text=Paragraph(el,d).text
                if i in figmap:result.append(f'![核验后结果图](../assets/20260928_{figmap[i]}.png)')
                elif el.xpath('.//a:blip'):
                    # Preserve unrelated four-class graphics through existing assets.
                    old={394:'fig5.6-2_four_class_prediction_and_discrimination.png',397:'fig5.6-3_top_label_reliability.png'}
                    if i in old:
                        candidate=r/'国赛报告/assets'/old[i]
                        assert candidate.is_file(),candidate
                        result.append(f'![四分类补充结果](../assets/{candidate.name})')
                elif text:
                    prefix='# ' if i==start else '## ' if re.match(r'^\d\.\d\.\d\s',text) else ''
                    result.append(prefix+text)
            elif el.tag==qn('w:tbl') and active:result.append(mdtable(Table(el,d)))
        return '\n\n'.join(result)+'\n'
    for start,end,name in [(359,403,'5.6-各科学模态对新参与者Q1的预测能力.md'),(403,427,'5.7-多模态增量、互补与条件预测价值.md'),(427,437,'5.8-设备组合与系统级输出.md')]:
        (r/'国赛报告/章节草稿'/name).write_text(extract(start,end),encoding='utf-8')
    method=r/'国赛报告/章节草稿/4.6-跨参与者监督学习与多模态比较.md';text=method.read_text(encoding='utf-8')
    def replace_prior(text,prefix,new):
        old=next((x for x in text.split('\n') if x.startswith(prefix) or x.startswith(new[:10])),None)
        if old is not None:return text.replace(old,new)
        assert new in text,prefix
        return text
    text=replace_prior(text,'首轮仅传感信息联合模型',paras[240].text)
    text=replace_prior(text,'信息类别层分别评价',paras[238].text)
    text=text.replace('随后获得正式预测资格的估计心率与估计呼吸率另组成 13 项指标的扩展组合，使用其自身的有效样本评价，不改写原 11 项主模型。','随后获得正式预测资格的估计心率与估计呼吸率组成 13 项完整系统，正文二分类完整组合、删减和校准统一采用该版本。保留先十一项、后十三项的分析顺序，原十一项四分类和精简验证仍按原定义报告。')
    method.write_text(text,encoding='utf-8')
    discussion=r/'国赛报告/章节草稿/6-分析与讨论.md';text=discussion.read_text(encoding='utf-8')
    text=replace_prior(text,'加入眼部后的',paras[522].text)
    text=replace_prior(text,'增加设备也伴随',paras[526].text);discussion.write_text(text,encoding='utf-8')
    method=r/'国赛报告/章节草稿/4.4-科学变量形成、窗口化与质量控制.md';text=method.read_text(encoding='utf-8').replace('心肺预测使用 110 场、2198 个探针的有效批次','心肺源表包含 110 场、2200 个有限值探针；两个独立单指标采用该集合，心肺类别及行为增量共同采用 110 场、2198 个探针');method.write_text(text,encoding='utf-8')
    appendix=r/'国赛报告/附录/第5章-眼部与预测完整结果附表.md';text=appendix.read_text(encoding='utf-8');lo=text.index('## B.2');hi=text.index('## B.4')
    text=text[:lo]+'## B.2 全部 28 项正式成对预测比较\n\n'+mdtable(d.tables[32])+'\n\n'+paras[659].text+'\n\n## B.3 全部正式模型的损失与排序诊断\n\n'+mdtable(d.tables[33])+'\n\n'+paras[662].text+'\n\n'+text[hi:];appendix.write_text(text,encoding='utf-8')
    (r/'国赛报告/附录/附录F-监督学习完整结果附表.md').write_text('# 附录 F 监督学习完整结果附表\n\n与附录 B.3 使用同一批 25 集合、56 行模型结果。95% 置信区间（confidence interval [CI]）采用固定折外预测的参与者簇自助法；原始数值及来源见正式结果总账。\n\n## F.1 各分析集合上的模型性能\n\n'+mdtable(d.tables[46])+'\n\n## F.2 概率诊断\n\n受试者工作特征曲线下面积（area under the receiver operating characteristic curve [AUROC]）仅对同时包含两类报告的参与者估计。布里尔分数（Brier score）是概率误差平方的参与者等权平均。判别力区间未排除 0.5 时，不作校准斜率的实质解释。\n\n'+mdtable(d.tables[47])+'\n',encoding='utf-8')
    evidence=r/'运行记录与证据'
    # Aggregate statistics only: never copy probe rows, participant identities or full logs.
    for name in ['model_results','paired_results','probability_diagnostics','session_results','constant_baselines']:
        src=a.results/(name+'.csv')
        import csv
        with src.open(encoding='utf-8-sig',newline='') as f:rd=csv.DictReader(f);fields=[x for x in rd.fieldnames if x not in ['source_run','run_id']];rows=[{k:v for k,v in x.items() if k in fields} for x in rd]
        with (evidence/('09-28-3-'+name+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:wr=csv.DictWriter(f,fieldnames=fields);wr.writeheader();wr.writerows(rows)
    audit=json.loads((a.results.parent/'audit_summary.json').read_text(encoding='utf-8'));independent=json.loads((a.results/'independent_verification.json').read_text(encoding='utf-8'));summary=json.loads((a.results/'manifest.json').read_text());edits=json.loads((a.results/'report_edit_manifest.json').read_text(encoding='utf-8'))
    safe={'status':'ANALYSIS_PASS_DOCUMENT_RENDER_PENDING','input_equivalence':audit['input_equivalence'],'historical_archive_folds':sum(x['outer_folds'] for x in audit['archive_audits']),'historical_inner_folds':sum(x['inner_folds'] for x in audit['archive_audits']),'result_summary':summary,'code_equivalence':independent['code_equivalence'],'report_input_sha256':edits['input_sha256'],'report_output_sha256':edits['output_sha256'],'changed_regions':edits['changed_regions'],'original_report_preserved':True,'replaced_figures':6}
    (evidence/'09-28-3-本机归档核验与最小重跑汇总.json').write_text(json.dumps(safe,ensure_ascii=False,indent=2),encoding='utf-8')
    print('Synced report sections, appendices, six figures and aggregate-only evidence')

if __name__=='__main__':main()
