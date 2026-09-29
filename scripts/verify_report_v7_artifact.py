"""Verify final package, current relationship targets, native captions and rendered PDF."""
import argparse,json,hashlib,re,zipfile,sys
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import pypdfium2 as pdfium
sys.stdout.reconfigure(encoding='utf-8')
ap=argparse.ArgumentParser();ap.add_argument('folder',type=Path);a=ap.parse_args();o=a.folder
sha=lambda b:hashlib.sha256(b).hexdigest()
m=json.loads((o/'edit_manifest.json').read_text(encoding='utf-8'));inv=json.loads((o/'resolved_image_inventory.json').read_text(encoding='utf-8'));plots=json.loads((o/'plot_evidence.json').read_text(encoding='utf-8'))
src=Path(m['input']);dst=Path(m['output']);assert sha(src.read_bytes())==m['input_sha256'];assert sha(dst.read_bytes())==m['output_sha256']
s=Document(src);d=Document(dst);ps=d.paragraphs
caps=[p for p in ps if p.style.name=='Caption'];assert len(caps)==86
for p in caps:
 assert p._p.xpath('./w:pPr/w:jc[@w:val="center"]')
 assert p._p.xpath('.//w:instrText[contains(text(),"SEQ")]')
 for r in p.runs:
  assert r._r.xpath('./w:rPr/w:sz[@w:val="18"]')
  assert r._r.xpath('./w:rPr/w:b[@w:val="1"]')
  assert r._r.xpath('./w:rPr/w:rFonts[@w:eastAsia="宋体"]')
figs=[p for p in caps if p.text.startswith('图 ')];tabs=[p for p in caps if p.text.startswith('表 ')];assert len(figs)==39 and len(tabs)==47
assert [int(re.match(r'图 (\d+)',p.text)[1]) for p in figs]==list(range(1,40))
assert sum(bool(p._p.xpath('.//w:drawing')) for p in ps)==39
plotmap={x['old_figure']:x for x in plots['figures']};checks=[]
for item in inv:
 old=item['figure']
 if str(old) not in m['figure_number_map']:continue
 new=m['figure_number_map'][str(old)];i=next(i for i,p in enumerate(ps) if p.text.startswith(f'图 {new} '))
 p=ps[i-1];blips=p._p.xpath('.//a:blip');assert len(blips)==1
 part=d.part.related_parts[blips[0].get(qn('r:embed'))];actual=sha(part.blob)
 assert not p._p.xpath('.//w:highlight[@w:val="yellow"]')
 if old in plotmap:expected=sha((o/plotmap[old]['png']).read_bytes());kind='source_based_statistical_plot'
 elif str(old) in m['original_clean_figures_restored']:
  asset=src.parent/'素材/系统界面图_20260913'/m['original_clean_figures_restored'][str(old)];expected=sha(asset.read_bytes());kind='original_clean_system_diagram'
 else:expected=item['sources'][0]['sha256'];kind='unchanged_original'
 assert actual==expected,(old,new,kind)
 checks.append({'old_figure':old,'new_figure':new,'type':kind,'sha256':actual,'caption':ps[i].text})
for t in d.tables[1:]:
 e=t._tbl.getprevious();assert e is not None and e.tag==qn('w:p')
 assert e.xpath('./w:pPr/w:pStyle')[0].get(qn('w:val'))==d.styles['Caption'].style_id
actual_changes=[]
for ti,(st,dt) in enumerate(zip(s.tables,d.tables)):
 for ri,(sr,dr) in enumerate(zip(st.rows,dt.rows)):
  for ci,(sc,dc) in enumerate(zip(sr.cells,dr.cells)):
   if sc.text!=dc.text:actual_changes.append((ti,ri,ci,dc.text))
expected_changes=[(x['table'],x['row'],x['column'],x['after']) for x in m['changes'] if 'table' in x]
assert sorted(actual_changes)==sorted(expected_changes)
assert [x._sectPr.xml for x in s.sections]==[x._sectPr.xml for x in d.sections]
with zipfile.ZipFile(src) as z1,zipfile.ZipFile(dst) as z2:
 assert z1.namelist()==z2.namelist()
 diff=[n for n in z1.namelist() if z1.read(n)!=z2.read(n)]
 assert set(diff)==set(m['replaced_parts']),diff
pdf=pdfium.PdfDocument(str(o/'v7_render.pdf'));pages=[p.get_textpage().get_text_range() for p in pdf]
assert all(t.strip() for t in pages)
assert not any(x in '\n'.join(pages) for x in ['Error! Reference source not found','错误!未找到引用源','错误！未找到引用源'])
norm=lambda x:re.sub(r'\s+','',x)
caption_pages={}
for p in caps:
 hits=[i+1 for i,t in enumerate(pages) if norm(p.text) in norm(t)]
 assert len(hits)==1,(p.text,hits)
 caption_pages[p.text]=hits[0]
result={'status':'PASS','docx_sha256':m['output_sha256'],'pdf_sha256':sha((o/'v7_render.pdf').read_bytes()),'pages':len(pdf),'native_captions':86,'figures':39,'tables_excluding_cover':47,'changed_table_cells':actual_changes,'figure_relationship_checks':checks,'caption_pages':caption_pages,'preserved_other_package_parts':True,'preserved_sections':True,'no_empty_pages':True,'models_refitted':0}
(o/'artifact_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print('PASS',len(pdf),'pages;',len(caps),'native captions;',len(checks),'image relationships')
