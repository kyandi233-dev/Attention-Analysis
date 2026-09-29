"""Preserve the report package; repair captions and replace source-based figures."""
import argparse,copy,hashlib,json,re,zipfile,sys
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
sys.stdout.reconfigure(encoding='utf-8')
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def text(p,value,yellow=True):
 for child in list(p._p):
  if child.tag!=qn('w:pPr'):p._p.remove(child)
 r=p.add_run(value)
 if yellow:r.font.highlight_color=7
def caption(p,kind,number,title,sequence=None,yellow=False):
 text(p,'',False);p.style='Caption'
 pp=p._p.get_or_add_pPr()
 for e in list(pp):
  if e.tag!=qn('w:pStyle'):pp.remove(e)
 for tag,val in [('jc','center'),('keepNext','1')]:
  e=OxmlElement('w:'+tag);e.set(qn('w:val'),val);pp.append(e)
 p.add_run(kind+' ')
 if sequence:p.add_run(sequence)
 for tag,attr,value in [('fldChar','fldCharType','begin'),('instrText',None,f' SEQ {kind} \\r {number} \\* ARABIC '),('fldChar','fldCharType','separate')]:
  e=OxmlElement('w:'+tag)
  if attr:e.set(qn('w:'+attr),value)
  else:e.text=value;e.set(qn('xml:space'),'preserve')
  p.add_run()._r.append(e)
 p.add_run(str(number));e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'end');p.add_run()._r.append(e)
 p.add_run(' '+title)
 for r in p.runs:
  rp=r._r.get_or_add_rPr();rf=OxmlElement('w:rFonts')
  for key,val in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia','宋体')]:rf.set(qn('w:'+key),val)
  rp.append(rf)
  for tag,val in [('sz','18'),('szCs','18'),('b','1'),('bCs','1')]:e=OxmlElement('w:'+tag);e.set(qn('w:val'),val);rp.append(e)
  if yellow:r.font.highlight_color=7
def main():
 ap=argparse.ArgumentParser();ap.add_argument('--source',type=Path,required=True);ap.add_argument('--results',type=Path,required=True);ap.add_argument('--assets',type=Path,required=True);a=ap.parse_args();out=a.results/'v7_audit';d=Document(a.source);ps=list(d.paragraphs);inventory=json.loads((out/'inventory.json').read_text(encoding='utf-8'));plots=json.loads((out/'plot_evidence.json').read_text(encoding='utf-8'));replacements={};changes=[]
 # Office may renumber package media parts when saving. Resolve every image
 # through the current document relationships; never reuse stale part names.
 for item in inventory:
  p=ps[item['paragraph']]
  assert re.match(r'^图\s*'+str(item['figure'])+r'\s',p.text),p.text
  drawings=ps[item['paragraph']-1]._p.xpath('.//a:blip')
  assert len(drawings)==1,(item['figure'],len(drawings))
  part=d.part.related_parts[drawings[0].get(qn('r:embed'))]
  item['sources']=[{'part':str(part.partname),'sha256':hashlib.sha256(part.blob).hexdigest()}]
 (out/'resolved_image_inventory.json').write_text(json.dumps(inventory,ensure_ascii=False,indent=2),encoding='utf-8')
 output=a.source.with_name('FocusWave_国赛报告_图表核验与题注统一_标黄副本_20260929_v7.docx')
 removed={8:'预实验形成性轨迹不占正式结果篇幅，正文保留方案形成过程',12:'反应时分布质量检查与方法说明重复',19:'质量热图与有效覆盖图重复',20:'视觉属性描述图移出，保留视觉控制方法',23:'眨眼描述轨迹与任务进程完整表重复',32:'问卷森林图与表21及附录A重复'}
 figuremap={old:i+1 for i,old in enumerate(n for n in range(1,46) if n not in removed)}
 def change(i,value):
  before=ps[i].text;text(ps[i],value);changes.append({'paragraph':i,'before':before,'after':value})
 def setcell(t,i,j,value):
  c=d.tables[t].rows[i].cells[j];p=c.paragraphs[0];before=c.text;rp=copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None;text(p,value)
  if rp is not None:
   run=p.runs[0];run._r.remove(run._r.rPr);run._r.insert(0,rp);run.font.highlight_color=7
  changes.append({'table':t,'row':i,'column':j,'before':before,'after':value})
 setcell(8,6,2,'探针前窗口内归一化画面变化率的中位数（s⁻¹）')
 setcell(34,13,2,'−0.004155');setcell(34,13,3,'0.01326');setcell(21,1,4,'.010')
 change(269,ps[269].text.replace('0.0078','0.0077').replace('[0.0033, 0.0122]','[0.0032, 0.0122]').replace('0.0001','0.0002').replace('[−0.0057, 0.0060]','[−0.0056, 0.0061]').replace('1391','1392'))
 change(277,'注：六项行为指标分别进入参与者聚类稳健多项逻辑回归，以任务聚焦为参照，均包含 61 名参与者、116 个场次，各指标有效探针数为 2316–2320。点为指标增加一个标准差时的对数优势变化，横线为 95% 置信区间。')
 change(286,ps[286].text+' 误差线为均值上下各 1.96 个标准误。')
 change(297,ps[297].text.replace('瞳孔对刺激尺寸、亮度与对比度的变化关系见图 20。',''))
 change(330,'注：点为整体身体动作强度的任务进程系数，横线为 95% 置信区间，虚线为零。模型按参与者聚类，包含 61 人、115 场和 1380 条阶段记录。区块项表示阶段编码为零时的区块差异，阶段项表示第一区块的变化，交互项表示两个区块的阶段斜率差。')
 change(434,'注：上方为任务诱发、多源采集和跨个体预测的研究流程，下方为状态反馈、单次总结和长期记录的交互流程。虚线框表示面向自然场景的后续验证环节。')
 from docx.text.paragraph import Paragraph
 method='单模态关联分析以参与者为重复观测的聚类单位。任务进程采用包含区块、区块内进程及其交互项的广义估计方程（generalized estimating equations [GEE]）。注意内容采用参与者聚类稳健多项逻辑回归；困倦—清醒等级在行为与动作分析中采用有序 GEE，在眼部分析中采用参与者聚类稳健有序逻辑回归。眼部模型同时纳入个人平均水平与相对个人平均的偏离，并控制区块和探针进程，按预先划定的分析族控制错误发现率。动作与行为的同期关系以动作指标为因变量，分别纳入标准化行为指标。'
 element=OxmlElement('w:p');ps[178]._p.addnext(element);p=Paragraph(element,d._body);p.style=ps[178].style;text(p,method);changes.append({'insert_after':178,'after':method})
 for idx,note in [(308,'注：点为优势比，横线为 95% 置信区间，虚线为无关联值 1。各类别均以任务聚焦为参照；优势比表示眼部指标的参与者内分量增加一个标准差时，相应报告的优势变化。'),(321,'注：上排为反应时变异系数与趋势的原尺度系数，下排为遗漏率与误按率的优势比。点与横线分别表示参与者内效应及 95% 置信区间；模型包含参与者平均水平、区块与探针进程。')]:
  element=OxmlElement('w:p');ps[idx]._p.addnext(element);p=Paragraph(element,d._body);p.style='注释';text(p,note)
 # Remove image paragraphs, their captions and notes. Keep all source files and source DOCX.
 for item in inventory:
  n=item['figure'];idx=item['paragraph']
  if n in removed:
   for j in [idx-1,idx,idx+1]:
    p=ps[j]
    if j==idx or (j==idx-1 and p._p.xpath('.//w:drawing')) or (j==idx+1 and p.text.startswith('注')):
     p._p.getparent().remove(p._p)
  else:
   title=re.sub(r'^图\s*\d+\s*','',item['caption'])
   title={15:'各模态指标的有效覆盖',16:'反应时变异系数的任务时间进程',17:'近期行为与注意内容报告的关系',25:'整体身体动作强度的任务进程效应'}.get(n,title)
   caption(ps[idx],'图',figuremap[n],title,yellow=(figuremap[n]!=n or n in [15,16,17,25]))
 # Six appendix graphics duplicate the complete ocular coefficient and coverage tables.
 start=next(p for p in ps if p.text=='B.4 眼部补充图')._p;end=next(p for p in ps if p.text=='B.5 数据来源')._p
 element=start
 while element is not end:
  following=element.getnext();element.getparent().remove(element);element=following
 p=next(p for p in ps if p._p is end);text(p,'B.4 数据来源')
 # Renumber prose references only after all old captions were reconstructed.
 for p in d.paragraphs:
  if p.style.name=='Caption':continue
  if re.search(r'图\s*\d+',p.text):
   value=re.sub(r'图\s*(\d+)',lambda m:'图 '+str(figuremap[int(m[1])]) if int(m[1]) in figuremap else m[0],p.text)
   if value!=p.text:text(p,value)
 # Normalize every main table caption, plus real native captions for all appendix tables.
 for p in d.paragraphs:
  match=re.match(r'^表\s*(\d+)\s*(.*)',p.text)
  if match:caption(p,'表',int(match[1]),match[2],yellow=bool(p._p.xpath('.//w:highlight[@w:val="yellow"]')))
 titles={22:('A',1,'事后问卷分析单位与样本'),23:('A',2,'事后走神等级分布'),24:('A',3,'事后问卷有序模型完整系数'),25:('B',1,'眼部任务进程的全部估计'),26:('B',2,'眼部与注意内容的参与者内关联'),27:('B',3,'眼部与注意内容的参与者间关联'),28:('B',4,'眼部与清醒等级的参与者内关联'),29:('B',5,'眼部与清醒等级的参与者间关联'),30:('B',6,'眼部与同期行为的参与者内关联'),31:('B',7,'眼部与同期行为的参与者间关联'),32:('B',8,'全部正式成对预测比较'),33:('B',9,'全部正式模型的损失与排序诊断'),34:('C',1,'两个正式区块之间的行为变化'),35:('C',2,'错误事件前后正确反应时的变化'),36:('C',3,'行为指标的有效覆盖'),37:('D',1,'动作指标的任务进程系数'),38:('D',2,'动作指标与即时注意内容的关系'),39:('D',3,'动作指标与清醒等级的关系'),40:('D',4,'动作指标与近期行为的关系'),41:('D',5,'动作指标的有效覆盖'),42:('E',1,'心率与呼吸率的有效覆盖'),43:('E',2,'心率估计误差的分类比较'),44:('E',3,'各类误差对总体低估的贡献'),45:('E',4,'不同心率估计路径的比较'),46:('F',1,'各分析集合上的模型性能'),47:('F',2,'概率诊断结果')}
 from docx.text.paragraph import Paragraph
 for idx,(letter,n,title) in titles.items():
  element=OxmlElement('w:p');d.tables[idx]._tbl.addprevious(element);caption(Paragraph(element,d._body),'表',n,title,sequence=letter,yellow=True)
 for item in plots['figures']:
  inv=next(x for x in inventory if x['figure']==item['old_figure']);replacements[inv['sources'][0]['part'].lstrip('/')]=(out/item['png']).read_bytes()
 restored={33:'系统图1_两条工作线与模型交接.png',34:'系统图2_测量结果的交接格式.png'}
 for n,name in restored.items():
  inv=next(x for x in inventory if x['figure']==n);replacements[inv['sources'][0]['part'].lstrip('/')]=(a.assets/name).read_bytes()
 # Resize only regenerated statistical drawings to their native aspect ratio.
 from PIL import Image
 for item in plots['figures']:
  n=item['old_figure'];p=ps[next(x for x in inventory if x['figure']==n)['paragraph']-1];w,h=Image.open(out/item['png']).size
  for ext in p._p.xpath('.//wp:extent'):
   width=int(ext.get('cx'));height=round(width*h/w);ext.set('cy',str(height))
   for ex in p._p.xpath('.//a:xfrm/a:ext'):ex.set('cy',str(height))
 for p in d.paragraphs:
  if p._p.xpath('.//w:drawing'):
   p.paragraph_format.keep_with_next=True
   for run in p.runs:
    if run._r.xpath('.//w:drawing'):run.font.highlight_color=None
 from lxml import etree
 replacements['word/document.xml']=etree.tostring(d._element,xml_declaration=True,encoding='UTF-8',standalone=True)
 with zipfile.ZipFile(a.source) as z,zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as dst:
  for info in z.infolist():dst.writestr(info,replacements.get(info.filename,z.read(info.filename)))
 check=Document(output);assert len(check.tables)==48
 captions=[p.text for p in check.paragraphs if p.style.name=='Caption'];assert len(captions)==86,(len(captions),captions)
 assert len([p for p in check.paragraphs if p._p.xpath('.//w:drawing')])==39
 assert [s._sectPr.xml for s in check.sections]==[s._sectPr.xml for s in Document(a.source).sections]
 manifest={'input':str(a.source),'input_sha256':sha(a.source),'output':str(output),'output_sha256':sha(output),'figure_number_map':figuremap,'removed_main_figures':removed,'removed_appendix_figures':6,'captions':86,'figures':39,'tables_excluding_cover':47,'regenerated_result_figures':14,'original_clean_figures_restored':restored,'replaced_parts':list(replacements),'changes':changes,'models_refitted':0}
 (out/'edit_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(manifest,ensure_ascii=False))
if __name__=='__main__':main()
