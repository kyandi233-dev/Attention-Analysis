"""Remove the separate attention-content four-class extension from the competition copy."""
import argparse,hashlib,json,re,zipfile
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

def replace_span(p,lo,hi,value):
    position=0;inserted=False
    for run in list(p.runs):
        end=position+len(run.text)
        if position<hi and end>lo:
            before=run.text[:max(0,lo-position)];after=run.text[max(0,hi-position):]
            run.text=before+(value if not inserted else '')+after
            run.font.highlight_color=7;inserted=True
        position=end
    assert inserted,(p.text,lo,hi)

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--results',type=Path,required=True);ap.add_argument('--formal-repo',type=Path);ap.add_argument('--verify',action='store_true');a=ap.parse_args();r=a.results
    m=json.loads((r/'v5_edit_manifest.json').read_text(encoding='utf-8'));source=Path(m['output']);output=source.with_name('FocusWave_国赛报告_主线精简_标黄副本_20260929_v6.docx')
    assert sha(source)==m['output_sha256']=='a17b526a2ded66f8ba664efb87acafcf1f3611b7eaa1adf8c7d3070f83e3da49'
    if a.verify:
        import pypdfium2 as pdfium
        from PIL import Image,ImageDraw
        old,d=Document(source),Document(output)
        assert len(d.tables)==48
        count=0
        for idx,t in enumerate(d.tables):
            for br,nr in zip(old.tables[idx].rows,t.rows):
                for bc,nc in zip(br.cells,nr.cells):assert bc.text==nc.text;count+=1
        text='\n'.join(p.text for p in d.paragraphs)
        assert text.count('四分类')==1 and 'RITnet 四分类语义分割' in text
        assert '附录 G' not in text and '顶端标签' not in text
        caps=[p for p in d.paragraphs if re.match(r'^[图表]\s+\d+\s',p.text)]
        assert len(caps)==66
        with zipfile.ZipFile(source) as z,zipfile.ZipFile(output) as y:
            assert z.namelist()==y.namelist()
            assert [n for n in z.namelist() if z.read(n)!=y.read(n)]==['word/document.xml']
        pdf=pdfium.PdfDocument(str(r/'v6_render.pdf'));pages=[p.get_textpage().get_text_range() for p in pdf];compact=re.sub(r'\s+','','\n'.join(pages))
        for p in caps:assert re.sub(r'\s+','',p.text)[:16] in compact,p.text
        (r/'v6_render_text.txt').write_text('\n\n'.join(f'PAGE {i+1}\n{p}' for i,p in enumerate(pages)),encoding='utf-8')
        dest=r/'render/v6';dest.mkdir(parents=True,exist_ok=True);thumbs=[]
        for i,page in enumerate(pdf):
            pic=page.render(scale=1.3).to_pil().convert('RGB');pic.save(dest/f'page-{i+1:03}.png');pic.thumbnail((265,365));tile=Image.new('RGB',(285,395),'white');tile.paste(pic,((285-pic.width)//2,23));ImageDraw.Draw(tile).text((8,5),str(i+1),fill='black');thumbs.append(tile)
        for offset in range(0,len(thumbs),12):
            pic=Image.new('RGB',(1140,1185),'#ccc')
            for j,t in enumerate(thumbs[offset:offset+12]):pic.paste(t,((j%4)*285,(j//4)*395))
            pic.save(dest/f'contact-{offset//12+1:02}.jpg')
        result={'status':'PASS','source_sha256':sha(source),'output_sha256':sha(output),'pages':len(pdf),'captions':66,'tables':48,'unchanged_table_cells':count,'models_refitted':0,'original_analysis_archives_preserved':True,'changed_package_parts':['word/document.xml']}
        (r/'v6_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(result));return
    if output.exists():assert sha(output)==json.loads((r/'v6_edit_manifest.json').read_text(encoding='utf-8'))['output_sha256']
    d=Document(source);ps=list(d.paragraphs);heading=next(p for p in ps if p.text=='5.6.5 四分类预测与补充分析');end=next(p for p in ps if p.text.startswith('头动协变量敏感性分析在'))
    removed=[];node=heading._p.getnext()
    while node is not end._p:
        nxt=node.getnext();removed.append(''.join(node.itertext()));node.getparent().remove(node);node=nxt
    replace_span(heading,0,len(heading.text),'5.6.5 头动协变量敏感性分析')
    appendix=next(p for p in d.paragraphs if p.text.startswith('附录 G'))
    node=appendix._p
    while node is not None and node.tag!=qn('w:sectPr'):
        nxt=node.getnext();node.getparent().remove(node);node=nxt
    for p in d.paragraphs:
        before='原 11 项组合的四分类与精简验证另作补充报告。'
        if before in p.text:
            at=p.text.index(before);replace_span(p,at,at+len(before),'原 11 项组合的精简验证见正文 5.7.3。')
        for match in reversed(list(re.finditer(r'图\s*(\d+)(?![\d.])',p.text))):
            n=int(match[1])
            if n>30:replace_span(p,match.start(1),match.end(1),str(n-2))
        for instr in p._p.xpath('.//w:instrText'):
            match=re.search(r'SEQ 图 \\r (\d+)',instr.text or '')
            if match and int(match[1])>30:instr.text=instr.text[:match.start(1)]+str(int(match[1])-2)+instr.text[match.end(1):]
    xml=etree.tostring(d._element,xml_declaration=True,encoding='UTF-8',standalone=True)
    with zipfile.ZipFile(source) as z,zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as y:
        for item in z.infolist():y.writestr(item,xml if item.filename=='word/document.xml' else z.read(item.filename))
    manifest={'source':str(source),'output':str(output),'source_sha256':sha(source),'output_sha256':sha(output),'removed_section':'Attention-content four-class extension only','removed_figures_v5':[29,30],'removed_appendix':'G','head_motion_sensitivity_retained':True,'models_refitted':0}
    (r/'v6_edit_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    if a.formal_repo:
        base=a.formal_repo/'国赛报告';f=base/'章节草稿/5.6-各科学模态对新参与者Q1的预测能力.md';s=f.read_text(encoding='utf-8')
        if '## 5.6.5 四分类预测与补充分析' not in s:
            print('Formal chapter already revised; numbering is not applied again.');return
        s=re.sub(r'## 5\.6\.5 四分类预测与补充分析.*?(?=头动协变量敏感性分析在)','## 5.6.5 头动协变量敏感性分析\n\n',s,flags=re.S);f.write_text(s,encoding='utf-8')
        for f in list((base/'章节草稿').glob('*.md'))+list((base/'附录').glob('*.md')):
            if f.name=='README.md':continue
            s=f.read_text(encoding='utf-8');old=s
            s=s.replace('原 11 项组合的四分类与精简验证另作补充报告。','原 11 项组合的精简验证见正文 5.7.3。').replace('窗口配对验证及四分类扩展','窗口配对验证及头动敏感性分析')
            s=re.sub(r'\n四类注意内容的分别预测已作为.*?(?=\n\n|\Z)','',s,flags=re.S)
            s=re.sub(r'(图\s*)(\d+)(?![\d.])',lambda m:m[1]+str(int(m[2])-2 if int(m[2])>30 else int(m[2])),s)
            if s!=old:f.write_text(s,encoding='utf-8')
    print(json.dumps(manifest,ensure_ascii=False))

if __name__=='__main__':main()
