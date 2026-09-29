"""Verify final layout edits without changing report evidence or caption styles."""
from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('manifest', type=Path)
    ap.add_argument('pdf', type=Path)
    ap.add_argument('output', type=Path)
    args = ap.parse_args()
    manifest = json.loads(args.manifest.read_text(encoding='utf-8'))
    source, target = Path(manifest['source']), Path(manifest['output'])
    assert sha256(source) == manifest['source_sha256']
    assert sha256(target) == manifest['output_sha256']
    with zipfile.ZipFile(source) as first, zipfile.ZipFile(target) as second:
        assert set(first.namelist()) == set(second.namelist())
        changed_parts = [name for name in first.namelist() if first.read(name) != second.read(name)]
        assert changed_parts == ['word/document.xml'], changed_parts
        media = [name for name in first.namelist() if name.startswith('word/media/')]
        assert media and all(first.read(name) == second.read(name) for name in media)
    before, after = Document(source), Document(target)
    assert len(before.paragraphs) == len(after.paragraphs) == 719
    assert len(before.tables) == len(after.tables) == 48
    assert [p.text for p in before.paragraphs] == [p.text for p in after.paragraphs]
    assert [[cell.text for row in t.rows for cell in row.cells] for t in before.tables] == [[cell.text for row in t.rows for cell in row.cells] for t in after.tables]
    caption_before = [p for p in before.paragraphs if p.style.name == 'Caption']
    caption_after = [p for p in after.paragraphs if p.style.name == 'Caption']
    assert len(caption_before) == len(caption_after) == 86
    assert all(etree.tostring(x._p) == etree.tostring(y._p) for x, y in zip(caption_before, caption_after))
    assert sum(p.text.startswith('图 ') for p in caption_after) == 39
    assert sum(p.text.startswith('表 ') for p in caption_after) == 47
    images = []
    for index, p in enumerate(after.paragraphs):
        if not p._p.xpath('.//w:drawing'):
            continue
        assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER, index
        assert not p.text and not p._p.xpath('.//w:t | .//w:tab | .//w:br'), index
        images.append(index)
    assert images == manifest['image_paragraphs_centered'] and len(images) == 39
    for index in manifest['ui_images_resized']:
        p = after.paragraphs[int(index)]
        assert len(p._p.xpath('.//wp:extent')) == 1
        assert round(int(p._p.xpath('.//wp:extent')[0].get('cy')) / 914400, 2) == 2.75
    assert not after.paragraphs[118]._p.xpath(".//w:br[@w:type='page']")
    assert len(after.paragraphs[606]._p.xpath('./w:pPr/w:sectPr')) == 1
    assert len(after.paragraphs[611]._p.xpath('./w:pPr/w:sectPr')) == 1
    assert not any(table._tbl.xpath('.//w:keepNext') for table in after.tables[22:])
    pdf = pdfium.PdfDocument(str(args.pdf))
    empty = [i + 1 for i, page in enumerate(pdf) if not page.get_textpage().get_text_range().strip()]
    assert len(pdf) == 103 and not empty
    result = {
        'status': 'PASS', 'source_sha256': sha256(source), 'output_sha256': sha256(target),
        'changed_package_parts': changed_parts, 'unchanged_media_files': len(media),
        'paragraphs': len(after.paragraphs), 'tables': len(after.tables),
        'figure_captions': 39, 'table_captions': 47, 'captions_changed': 0,
        'text_changes': 0, 'table_value_changes': 0, 'figure_content_changes': 0,
        'images_centered_without_text_tabs_breaks': len(images),
        'ui_images_resized': len(manifest['ui_images_resized']),
        'rendered_pages': len(pdf), 'empty_pages': empty,
        'landscape_appendix_sections_preserved': True,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'status': result['status'], 'pages': result['rendered_pages'], 'figures': len(images), 'captions': 86}, ensure_ascii=False))


if __name__ == '__main__':
    main()
