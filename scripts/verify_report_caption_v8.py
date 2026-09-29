"""Check every restored Word caption against the source style and visible PDF."""
import argparse,hashlib,json,re,zipfile
from pathlib import Path
import pypdfium2 as pdfium
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
ap=argparse.ArgumentParser();ap.add_argument('folder',type=Path);a=ap.parse_args();o=a.folder
m=json.loads((o/'caption_v8_manifest.json').read_text(encoding='utf-8'));p7=Path(m['input']);p8=Path(m['output']);ref=Document(m['reference']);d7=Document(p7);d=Document(p8)
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
assert sha(p8)==m['output_sha256'] and sha(p7)==m['input_sha256']
with zipfile.ZipFile(p7) as a,zipfile.ZipFile(p8) as b:
 assert a.namelist()==b.namelist()
 assert [x for x in a.namelist() if a.read(x)!=b.read(x)]==['word/document.xml']
body7=[p.text for p in d7.paragraphs if p.style.name!='Caption'];body8=[p.text for p in d.paragraphs if p.style.name!='Caption'];assert len(body7)==len(body8)
changed=[(x,y) for x,y in zip(body7,body8) if x!=y];assert len(changed)==2
assert len(d.tables)==len(d7.tables)==48
for t,u in zip(d.tables,d7.tables):
 assert [[c.text for c in r.cells] for r in t.rows]==[[c.text for c in r.cells] for r in u.rows]
cs=[p for p in d.paragraphs if p.style.name=='Caption'];assert len(cs)==86
oldmap=json.loads((o/'edit_manifest.json').read_text(encoding='utf-8'))['figure_number_map'];inverse={int(v):int(k) for k,v in oldmap.items()}
ref_fig={int(re.match(r'图\s*(\d+)',p.text)[1]):p for p in ref.paragraphs if p.style.name=='Caption' and re.match(r'^图\s*\d+',p.text)}
ref_tab={int(re.match(r'表\s*(\d+)',p.text)[1]):p for p in ref.paragraphs if p.style.name=='Caption' and re.match(r'^表\s*\d+',p.text)}
assert len([p for p in cs if p.text.startswith('图 ')])==39
assert len([p for p in cs if p.text.startswith('表 ')])==47
assert [int(re.match(r'图 (\d+)',p.text)[1]) for p in cs if p.text.startswith('图 ')]==list(range(1,40))
for p in cs:
 assert p._p.xpath('.//w:instrText[contains(text(),"SEQ")]'),p.text
 assert p._p.pPr is not None and p._p.pPr.xpath('./w:pStyle')[0].get(qn('w:val'))==d.styles['Caption'].style_id
 label=re.match(r'^(图|表)\s*([A-F]?)(\d+)',p.text);kind,letter,num=label.groups();num=int(num)
 reference=ref_fig[inverse[num]] if kind=='图' else ref_tab[1 if letter else num]
 assert etree.tostring(p._p.pPr)==etree.tostring(reference._p.pPr),p.text
 for r in p.runs:
  if r._r.rPr is not None:assert all(child.tag==qn('w:highlight') for child in r._r.rPr),p.text
for i,p in enumerate(d.paragraphs):
 if p.style.name=='Caption' and p.text.startswith('图 '):assert i>0 and d.paragraphs[i-1]._p.xpath('.//w:drawing'),p.text
for t in d.tables[1:]:
 prev=t._tbl.getprevious();assert prev is not None and prev.tag==qn('w:p')
 assert prev.xpath('./w:pPr/w:pStyle')[0].get(qn('w:val'))==d.styles['Caption'].style_id
assert sum(bool(p._p.xpath('.//w:drawing')) for p in d.paragraphs)==39
pdf=pdfium.PdfDocument(str(o/'v8_render.pdf'));texts=[p.get_textpage().get_text_range() for p in pdf];assert len(pdf)>0 and all(x.strip() for x in texts)
alltext=''.join(texts);norm=lambda s:re.sub(r'\s+','',s)
for p in cs:assert sum(norm(p.text) in norm(s) for s in texts)==1,p.text
assert '结果见5.6.5' not in norm(alltext)
assert not any(x in alltext for x in ['Error! Reference source not found','错误!未找到引用源','错误！未找到引用源'])
report={'status':'PASS','source_report_sha256':m['input_sha256'],'repaired_report_sha256':m['output_sha256'],'pages':len(pdf),'native_word_captions':86,'figures':39,'tables':47,'uncaptioned_figures':0,'uncaptioned_tables':0,'media_parts_changed':0,'table_values_changed':0,'other_paragraphs_changed':len(changed),'changes':changed,'original_caption_style_id':ref.styles['Caption'].style_id,'output_caption_style_id':d.styles['Caption'].style_id}
(o/'caption_v8_verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');print('PASS',len(pdf),'pages; 86 native captions; all tables and figures paired')
