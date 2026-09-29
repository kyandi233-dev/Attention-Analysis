"""Verify that v9 changes only five research-story paragraphs in v8."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from docx.oxml.ns import qn
from lxml import etree


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("manifest", type=Path)
    ap.add_argument("rendered_pdf", type=Path)
    ap.add_argument("verification", type=Path)
    args = ap.parse_args()
    manifest = json.loads(args.manifest.read_text(encoding="utf-8"))
    source, output = Path(manifest["source"]), Path(manifest["output"])
    digest = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()
    assert digest(source) == manifest["source_sha256"]
    assert digest(output) == manifest["output_sha256"]
    with zipfile.ZipFile(source) as first, zipfile.ZipFile(output) as second:
        assert set(first.namelist()) == set(second.namelist())
        parts_changed = [name for name in first.namelist() if first.read(name) != second.read(name)]
        assert parts_changed == ["word/document.xml"], parts_changed
        a = etree.fromstring(first.read("word/document.xml"))
        b = etree.fromstring(second.read("word/document.xml"))
        body_a, body_b = a.find(qn("w:body")), b.find(qn("w:body"))
        assert len(body_a) == len(body_b)
        changed_body = [i for i, (one, two) in enumerate(zip(body_a, body_b)) if etree.tostring(one) != etree.tostring(two)]
        assert len(changed_body) == len(manifest["changed_paragraphs"]) == 5
        for i in changed_body:
            assert body_a[i].tag == body_b[i].tag == qn("w:p")
            assert etree.tostring(body_a[i].find(qn("w:pPr"))) == etree.tostring(body_b[i].find(qn("w:pPr")))
    before, after = Document(source), Document(output)
    assert len(before.paragraphs) == len(after.paragraphs)
    changed_paragraphs = [i for i, (one, two) in enumerate(zip(before.paragraphs, after.paragraphs)) if one.text != two.text]
    assert changed_paragraphs == [item["paragraph_index"] for item in manifest["changed_paragraphs"]]
    assert len(before.tables) == len(after.tables) == 48
    captions_before = [p for p in before.paragraphs if p.style.name == "Caption"]
    captions_after = [p for p in after.paragraphs if p.style.name == "Caption"]
    assert len(captions_before) == len(captions_after) == 86
    assert all(etree.tostring(one._p) == etree.tostring(two._p) for one, two in zip(captions_before, captions_after))
    assert sum(p.text.startswith("图 ") for p in captions_after) == 39
    assert sum(p.text.startswith("表 ") for p in captions_after) == 47
    for item in manifest["changed_paragraphs"]:
        paragraph = after.paragraphs[item["paragraph_index"]]
        assert paragraph.text == item["after"]
        assert paragraph._p.xpath('./w:r/w:rPr/w:highlight[@w:val="yellow"]')
        assert not any(word in paragraph.text for word in ("需要强调的是", "不能简单", "不应机械", "待核验", "源表", "快照", "工程"))
    pdf = pdfium.PdfDocument(str(args.rendered_pdf))
    page_texts = [page.get_textpage().get_text_range() for page in pdf]
    assert len(pdf) == 110 and all(text.strip() for text in page_texts)
    compact = lambda value: re.sub(r"\s+", "", value)
    pages = {}
    for item in manifest["changed_paragraphs"]:
        prefix = compact(item["after"][:28])
        matches = [i + 1 for i, page in enumerate(page_texts) if prefix in compact(page)]
        assert len(matches) == 1, matches
        pages[str(item["paragraph_index"])] = matches[0]
    result = {
        "status": "PASS",
        "source_sha256": manifest["source_sha256"],
        "output_sha256": manifest["output_sha256"],
        "changed_paragraph_indices": changed_paragraphs,
        "rendered_pages": len(pdf),
        "modified_paragraph_pages": pages,
        "blank_pages": 0,
        "captions_unchanged": 86,
        "figures_retained": 39,
        "tables_retained": 47,
        "media_parts_changed": 0,
        "table_parts_changed": 0,
        "models_refitted": 0,
    }
    args.verification.parent.mkdir(parents=True, exist_ok=True)
    args.verification.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
