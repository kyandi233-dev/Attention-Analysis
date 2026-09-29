"""Check the competition-style revision against its sources and native rendering."""
import argparse, hashlib, json, re, zipfile
from pathlib import Path
from docx import Document
import pypdfium2 as pdfium
from PIL import Image, ImageDraw

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--results',type=Path,required=True);a=ap.parse_args();root=a.results
    m=json.loads((root/'v5_edit_manifest.json').read_text(encoding='utf-8'))
    sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
    assert sha(m['input'])==m['input_sha256']
    assert sha(m['output'])==m['output_sha256']
    original=Path(m['input']).parent/'FocusWave_国赛报告_润色副本.docx'
    assert sha(original)==m['original_sha256']
    with zipfile.ZipFile(m['input']) as x,zipfile.ZipFile(m['output']) as y:
        assert x.namelist()==y.namelist()
        changed=[n for n in x.namelist() if x.read(n)!=y.read(n)]
        assert sorted(changed)==sorted(['word/document.xml']+m['replaced_media'])
    old,new,orig=Document(m['input']),Document(m['output']),Document(original)
    assert len(old.tables)==len(new.tables)==51
    assert [s._sectPr.xml for s in old.sections]==[s._sectPr.xml for s in new.sections]
    original_captions={}
    for p in orig.paragraphs:
        match=re.match(r'^([图表])\s+(\d+)\s',p.text)
        if match:original_captions[(match[1],int(match[2]))]=p
    captions=[]
    for p in new.paragraphs:
        match=re.match(r'^([图表])\s+(\d+)\s',p.text)
        if not match:continue
        kind,n=match[1],int(match[2]);source=original_captions[(kind,n+1 if kind=='图' and n>=9 else n)]
        assert p._p.pPr.xml==source._p.pPr.xml,(kind,n)
        assert any('SEQ' in (v.text or '') for v in p._p.xpath('.//w:instrText'))
        captions.append(p.text)
    assert len(captions)==68
    text='\n'.join(p.text for p in new.paragraphs)+'\n'+'\n'.join(c.text for t in new.tables for r in t.rows for c in r.cells)
    for forbidden in ['无已检测时序歧义','旧测量快照','生产版本','本图仅用于形成性描述','不构成独立参与者','2198 个有效心肺','2180','内部质量指标的描述统计']:
        assert forbidden not in text,forbidden
    count=0
    for idx in [14,15,18,19,20,32,33,37,38,39,40,41,46,47,48,49,50]:
        b=[c.text for r in old.tables[idx].rows for c in r.cells];n=[c.text for r in new.tables[idx].rows for c in r.cells]
        assert len(b)==len(n)
        for before,after in zip(b,n):
            if re.fullmatch(r'[\d\s.,%−+\-\[\]()/<>≤≥=NAna—–]+',before) and re.search(r'\d',before):
                assert before==after,(idx,before,after);count+=1
    payload=json.loads((root/'report_payload.json').read_text(encoding='utf-8'))
    fresh=json.loads((root/'v5_plot_payload.json').read_text(encoding='utf-8'))
    assert payload['models']==fresh['models'] and payload['pairs']==fresh['pairs']
    fmt=lambda v,n=4:f'{v:.{n}f}'.replace('-','−')
    for raw,row in zip(payload['models'],new.tables[46].rows[1:]):
        assert [c.text for c in row.cells][2:7]==[str(raw['n_participants']),str(raw['n_probes']),fmt(raw['participant_equal_log_loss']),fmt(raw['ci_lower']),fmt(raw['ci_upper'])]
    for raw,row in zip(payload['pairs'],new.tables[32].rows[1:]):
        assert row.cells[2].text==fmt(raw['point_estimate'],6)+' ['+fmt(raw['ci_lower'],6)+', '+fmt(raw['ci_upper'],6)+']'
    assert len(new.tables[13].rows)==3 and len(new.tables[13].columns)==6
    assert all(row.cells[1].text=='2200' for row in new.tables[13].rows[1:])
    pdf=pdfium.PdfDocument(str(root/'v5_render.pdf'));pages=[p.get_textpage().get_text_range() for p in pdf]
    compact=re.sub(r'\s+','','\n'.join(pages))
    for caption in captions:assert re.sub(r'\s+','',caption)[:16] in compact,caption
    (root/'v5_render_text.txt').write_text('\n\n'.join(f'PAGE {i+1}\n{t}' for i,t in enumerate(pages)),encoding='utf-8')
    out=root/'render/v5';out.mkdir(parents=True,exist_ok=True);thumbs=[]
    for i,page in enumerate(pdf):
        pic=page.render(scale=1.3).to_pil().convert('RGB');pic.save(out/f'page-{i+1:03}.png');pic.thumbnail((265,365))
        tile=Image.new('RGB',(285,395),'white');tile.paste(pic,((285-pic.width)//2,23));ImageDraw.Draw(tile).text((8,5),f'Page {i+1}',fill='black');thumbs.append(tile)
    for offset in range(0,len(thumbs),12):
        sheet=Image.new('RGB',(1140,1185),'#c0c0c0')
        for j,pic in enumerate(thumbs[offset:offset+12]):sheet.paste(pic,((j%4)*285,(j//4)*395))
        sheet.save(out/f'contact-{offset//12+1:02}.jpg')
    result={'status':'PASS','output_sha256':sha(m['output']),'original_unchanged':True,'source_v4_unchanged':True,'changed_package_parts':changed,'styles_headers_footers_unchanged':True,'caption_paragraph_formats_restored':68,'native_caption_numbers_matched':68,'preserved_numeric_cells':count,'frozen_model_rows':len(payload['models']),'frozen_pair_rows':len(payload['pairs']),'predictive_refits':0,'pages':len(pdf),'pages_without_extractable_text':[i+1 for i,t in enumerate(pages) if not t.strip()],'visual_review':'See run record; separate contact-sheet and enlarged-page inspection.'}
    (root/'v5_document_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(result,ensure_ascii=False))

if __name__=='__main__':main()
