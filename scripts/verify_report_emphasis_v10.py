"""Verify full-report emphasis review and v9-to-v10 formatting-only delta."""
from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from docx.oxml.ns import qn
from lxml import etree


def positions(paragraph, attribute):
    mask = []
    for run in paragraph.runs:
        mask.extend([bool(getattr(run, attribute))] * len(run.text))
    assert len(mask) == len(paragraph.text)
    return mask


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('manifest', type=Path)
    parser.add_argument('pdf', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    manifest = json.loads(args.manifest.read_text(encoding='utf-8'))
    a, b = Path(manifest['source']), Path(manifest['output'])
    sha = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()
    assert sha(a) == manifest['source_sha256'] and sha(b) == manifest['output_sha256']
    with zipfile.ZipFile(a) as first, zipfile.ZipFile(b) as second:
        assert set(first.namelist()) == set(second.namelist())
        changed_parts = [name for name in first.namelist() if first.read(name) != second.read(name)]
        assert changed_parts == ['word/document.xml']
        before_xml = etree.fromstring(first.read('word/document.xml'))
        after_xml = etree.fromstring(second.read('word/document.xml'))
        before_body = before_xml.find(qn('w:body'))
        after_body = after_xml.find(qn('w:body'))
        assert len(before_body) == len(after_body)
        changed_elements = [i for i, (x, y) in enumerate(zip(before_body, after_body)) if etree.tostring(x) != etree.tostring(y)]
        assert len(changed_elements) == len(manifest['restyled_paragraphs'])
        for i in changed_elements:
            x, y = before_body[i], after_body[i]
            assert x.tag == y.tag == qn('w:p')
            assert etree.tostring(x.find(qn('w:pPr'))) == etree.tostring(y.find(qn('w:pPr')))
    before, after = Document(a), Document(b)
    assert len(before.paragraphs) == len(after.paragraphs)
    assert [p.text for p in before.paragraphs] == [p.text for p in after.paragraphs]
    assert len(before.tables) == len(after.tables) == 48
    changed_paragraphs = [i for i, (x, y) in enumerate(zip(before.paragraphs, after.paragraphs)) if etree.tostring(x._p) != etree.tostring(y._p)]
    assert changed_paragraphs == [int(i) for i in manifest['restyled_paragraphs']]
    before_captions = [p for p in before.paragraphs if p.style.name == 'Caption']
    after_captions = [p for p in after.paragraphs if p.style.name == 'Caption']
    assert len(before_captions) == len(after_captions) == 86
    assert all(etree.tostring(x._p) == etree.tostring(y._p) for x, y in zip(before_captions, after_captions))
    assert sum(p.text.startswith('图 ') for p in after_captions) == 39
    assert sum(p.text.startswith('表 ') for p in after_captions) == 47
    for index, plan in manifest['restyled_paragraphs'].items():
        p = after.paragraphs[int(index)]
        bold, underline = positions(p, 'bold'), positions(p, 'underline')
        assert not any(x and y for x, y in zip(bold, underline))
        expected_bold = [False] * len(p.text)
        expected_under = [False] * len(p.text)
        for phrase in plan['bold']:
            assert p.text.count(phrase) == 1
            start = p.text.index(phrase)
            expected_bold[start:start + len(phrase)] = [True] * len(phrase)
        for phrase in plan['underline']:
            assert p.text.count(phrase) == 1
            start = p.text.index(phrase)
            expected_under[start:start + len(phrase)] = [True] * len(phrase)
        assert bold == expected_bold, index
        assert underline == expected_under, index
    full_document_dual = [(i, r.text) for i, p in enumerate(after.paragraphs) for r in p.runs if r.bold and r.underline and r.text.strip()]
    assert not full_document_dual
    table_dual = [(ti, ri, ci) for ti, table in enumerate(after.tables) for ri, row in enumerate(table.rows) for ci, cell in enumerate(row.cells) for p in cell.paragraphs for r in p.runs if r.bold and r.underline and r.text.strip()]
    assert not table_dual
    pdf = pdfium.PdfDocument(str(args.pdf))
    blank = [i + 1 for i, page in enumerate(pdf) if not page.get_textpage().get_text_range().strip()]
    assert len(pdf) == 110 and not blank
    report = {'status': 'PASS', 'source_sha256': manifest['source_sha256'], 'output_sha256': manifest['output_sha256'], 'paragraphs_inspected': len(after.paragraphs), 'tables_inspected': len(after.tables), 'emphasis_paragraphs_repaired': changed_paragraphs, 'text_changes': 0, 'caption_changes': 0, 'figure_changes': 0, 'table_changes': 0, 'full_document_dual_emphasis': 0, 'table_dual_emphasis': 0, 'rendered_pages': len(pdf), 'blank_pages': blank}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'status': 'PASS', 'repaired': len(changed_paragraphs), 'pages': len(pdf)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
