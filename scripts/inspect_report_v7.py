"""Read-only DOCX inventory for the figure/caption review."""
import argparse,json,re,sys
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
ap=argparse.ArgumentParser();ap.add_argument('source',type=Path);ap.add_argument('out',type=Path);ap.add_argument('--prefix',default='');a=ap.parse_args();d=Document(a.source)
ps={id(p._p):i for i,p in enumerate(d.paragraphs)}
hist=[];tables=[];index=0
for element in d._element.body:
 if element.tag==qn('w:p'):
  text=''.join(element.itertext()) if False else ''.join(element.xpath('.//w:t/text()'))
  if text.strip():hist.append(text)
 elif element.tag==qn('w:tbl'):
  t=d.tables[index];tables.append({'index':index,'before':hist[-3:],'rows':[[c.text for c in row.cells] for row in t.rows]});index+=1
(a.out/(a.prefix+'table_inventory.json')).write_text(json.dumps(tables,ensure_ascii=False,indent=2),encoding='utf-8')
(a.out/(a.prefix+'paragraph_inventory.txt')).write_text('\n'.join(f'{i}\t{p.style.name}\t{p.text}' for i,p in enumerate(d.paragraphs)),encoding='utf-8')
for t in tables:print(t['index'],len(t['rows']),' | '.join(t['before'][-1:])[:100],t['rows'][0], t['rows'][1] if len(t['rows'])>1 else [])
