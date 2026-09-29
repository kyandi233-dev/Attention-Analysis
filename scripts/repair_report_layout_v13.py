"""Repair avoidable report whitespace while preserving text and figure content."""
from __future__ import annotations

import argparse
import hashlib
import json
import io
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("manifest", type=Path)
    parser.add_argument("--ui-height-inches", type=float, default=3.25)
    parser.add_argument("--f2-cell-space-pt", type=float)
    args = parser.parse_args()
    assert args.source.is_file() and not args.output.exists()
    doc = Document(args.source)
    assert len(doc.paragraphs) == 719 and len(doc.tables) == 48

    image_indices = []
    resized = {}
    for index, paragraph in enumerate(doc.paragraphs):
        drawings = paragraph._p.xpath(".//w:drawing")
        if not drawings:
            continue
        assert len(drawings) == 1
        assert not paragraph.text.strip()
        assert not paragraph._p.xpath(".//w:t | .//w:tab | .//w:br")
        image_indices.append(index)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = 0
        paragraph.paragraph_format.left_indent = 0
        paragraph.paragraph_format.right_indent = 0
        if 428 <= index <= 460:
            extents = paragraph._p.xpath(".//wp:extent")
            assert len(extents) == 1
            extent = extents[0]
            old_cx, old_cy = int(extent.get("cx")), int(extent.get("cy"))
            # Keep mockup proportions; a slightly smaller print size allows
            # the neighboring caption and prose to share a page.
            max_height = int(args.ui_height_inches * 914400)
            if old_cy > max_height:
                ratio = max_height / old_cy
                extent.set("cx", str(round(old_cx * ratio)))
                extent.set("cy", str(max_height))
                for transform in paragraph._p.xpath(".//a:xfrm/a:ext"):
                    transform.set("cx", str(round(old_cx * ratio)))
                    transform.set("cy", str(max_height))
                resized[index] = {
                    "old_inches": [round(old_cx / 914400, 3), round(old_cy / 914400, 3)],
                    "new_inches": [round(old_cx * ratio / 914400, 3), args.ui_height_inches],
                }

    assert len(image_indices) == 39
    assert len(resized) == 9
    # The page break following the short section 3.3 introduction created
    # two successive sparse pages before the first method figures.
    page_breaks = doc.paragraphs[118]._p.xpath(".//w:br[@w:type='page']")
    assert len(page_breaks) == 1
    page_breaks[0].getparent().remove(page_breaks[0])

    cleared = 0
    for table in doc.tables[22:]:
        for keep in table._tbl.xpath(".//w:keepNext"):
            keep.getparent().remove(keep)
            cleared += 1
    assert cleared > 0
    if args.f2_cell_space_pt is not None:
        for row in doc.tables[47].rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(args.f2_cell_space_pt)
                    paragraph.paragraph_format.space_after = Pt(args.f2_cell_space_pt)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    buffer = io.BytesIO()
    doc.save(buffer)
    with zipfile.ZipFile(args.source) as original, zipfile.ZipFile(buffer) as revised, zipfile.ZipFile(args.output, 'w') as final:
        revised_document = revised.read('word/document.xml')
        for item in original.infolist():
            final.writestr(item, revised_document if item.filename == 'word/document.xml' else original.read(item.filename))
    manifest = {
        "source": str(args.source),
        "source_sha256": sha256(args.source),
        "output": str(args.output),
        "output_sha256": sha256(args.output),
        "image_paragraphs_centered": image_indices,
        "image_paragraphs_without_text_tabs_breaks": len(image_indices),
        "ui_images_resized": resized,
        "ui_height_inches": args.ui_height_inches,
        "removed_page_break_after_paragraph": 118,
        "appendix_table_keep_with_next_cleared": cleared,
        "f2_cell_space_pt": args.f2_cell_space_pt,
    }
    args.manifest.parent.mkdir(parents=True, exist_ok=True)
    args.manifest.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": "BUILT", "figures": len(image_indices), "resized": len(resized), "table_keep_cleared": cleared}, ensure_ascii=False))


if __name__ == "__main__":
    main()
